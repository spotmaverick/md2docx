# -*- coding: utf-8 -*-
"""AST -> DOCX（原生 Word 2007+ 格式，WPS 可直接打开）。"""
from __future__ import annotations

import io
import os
import urllib.request
from urllib.parse import unquote, urlparse

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Emu, Pt, RGBColor

import theme
from mdmodel import Block, Run

UA = {"User-Agent": "Mozilla/5.0 Md2docs"}


# --------------------------------------------------------------------------- #
# 底层小工具
# --------------------------------------------------------------------------- #
def _sub(parent, tag, **attrs):
    el = OxmlElement(tag)
    for k, v in attrs.items():
        el.set(qn("w:" + k), str(v))
    parent.append(el)
    return el


def _shd(parent, fill: str):
    _sub(parent, "w:shd", val="clear", color="auto", fill=fill)


def _font_el(run):
    rPr = run._element.get_or_add_rPr()
    fonts = rPr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rPr.insert(0, fonts)
    return fonts


def style_run(run, name: str | None = None, ea: str | None = None,
              size: float | None = None, bold: bool | None = None,
              italic: bool | None = None, color: str | None = None,
              mono: bool = False, underline: bool | None = None,
              strike: bool | None = None, sup: bool = False, sub: bool = False,
              shading: str | None = None):
    if mono:
        name, ea = theme.MONO_FONT, theme.MONO_FONT
    fonts = _font_el(run)
    if name:
        fonts.set(qn("w:ascii"), name)
        fonts.set(qn("w:hAnsi"), name)
    if ea:
        fonts.set(qn("w:eastAsia"), ea)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if underline is not None:
        run.underline = underline
    if strike is not None:
        run.font.strike = strike
    if sup or sub:
        run.font.superscript = sup
        run.font.subscript = sub
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    if shading:
        _shd(run._element.get_or_add_rPr(), shading)


def add_hyperlink(paragraph, url: str, text: str):
    part = paragraph.part
    r_id = part.relate_to(url,
                          "http://schemas.openxmlformats.org/officeDocument/"
                          "2006/relationships/hyperlink",
                          is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), r_id)
    new_run = OxmlElement("w:r")
    rPr = OxmlElement("w:rPr")
    new_run.append(rPr)
    r = _sub(rPr, "w:rStyle", val="Hyperlink") if False else None
    fonts = OxmlElement("w:rFonts")
    fonts.set(qn("w:ascii"), theme.BODY_FONT)
    fonts.set(qn("w:hAnsi"), theme.BODY_FONT)
    fonts.set(qn("w:eastAsia"), theme.BODY_FONT_EA)
    rPr.append(fonts)
    _sub(rPr, "w:color", val=theme.LINK)
    _sub(rPr, "w:u", val="single")
    _sub(rPr, "w:sz", val="22")
    t = OxmlElement("w:t")
    t.text = text
    t.set(qn("xml:space"), "preserve")
    new_run.append(t)
    link.append(new_run)
    paragraph._p.append(link)
    return link


# --------------------------------------------------------------------------- #
# 资源加载（相对路径 / 网络图片）
# --------------------------------------------------------------------------- #
def load_image(src: str, base_dir: str) -> tuple[bytes, str] | None:
    if not src:
        return None
    try:
        if src.startswith(("http://", "https://", "//")):
            url = ("https:" + src) if src.startswith("//") else src
            with urllib.request.urlopen(url, timeout=15) as f:
                data = f.read()
            ext = os.path.splitext(urlparse(url).path)[1].lower() or ".png"
            return data, ext
        if src.startswith("data:"):
            import base64
            head, payload = src.split(",", 1)
            data = base64.b64decode(payload)
            ext = ".png"
            if "jpeg" in head or "jpg" in head:
                ext = ".jpg"
            elif "gif" in head:
                ext = ".gif"
            return data, ext
        path = unquote(src.replace("/", os.sep))
        if not os.path.isabs(path):
            path = os.path.join(base_dir, path)
        if os.path.isfile(path):
            with open(path, "rb") as f:
                data = f.read()
            ext = os.path.splitext(path)[1].lower() or ".png"
            return data, ext
    except Exception:
        return None
    return None


# --------------------------------------------------------------------------- #
# 主入口
# --------------------------------------------------------------------------- #
def write_docx(blocks: list[Block], out_path: str, base_dir: str = "",
               title: str = "") -> None:
    doc = Document()

    # 页面：A4 + 2.5cm 边距
    for sec in doc.sections:
        sec.page_width = Cm(21.0)
        sec.page_height = Cm(29.7)
        sec.left_margin = Cm(2.5)
        sec.right_margin = Cm(2.5)
        sec.top_margin = Cm(2.5)
        sec.bottom_margin = Cm(2.5)

    normal = doc.styles["Normal"]
    normal.font.name = theme.BODY_FONT
    normal.font.size = Pt(theme.BODY_SIZE)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), theme.BODY_FONT_EA)
    pf = normal.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(6)
    pf.line_spacing = 1.5

    for lvl in range(1, 7):
        try:
            st = doc.styles[f"Heading {lvl}"]
        except KeyError:
            continue
        # 与 Markdown 渲染一致：标题不额外着色，仅放大字号 + 加粗
        st.font.name = theme.BODY_FONT
        st.font.size = Pt(theme.heading_size(lvl))
        st.font.bold = True
        st.font.color.rgb = RGBColor.from_string(theme.HEADING_TEXT)
        st.element.rPr.rFonts.set(qn("w:eastAsia"), theme.BODY_FONT_EA)
        st.paragraph_format.space_before = Pt(14 if lvl <= 2 else 8)
        st.paragraph_format.space_after = Pt(6)
        st.paragraph_format.keep_with_next = True

    if title:
        doc.core_properties.title = title

    body = doc
    for b in blocks:
        _render_block(body, b, base_dir)

    doc.save(out_path)


def _content_width(doc) -> Emu:
    sec = doc.sections[0]
    return sec.page_width - sec.left_margin - sec.right_margin


def _render_block(doc, b: Block, base_dir: str, **ctx):
    if b.kind == "heading":
        p = doc.add_paragraph(style=f"Heading {min(b.level, 6)}")
        _emit_runs(p, b.runs, base_dir, size=theme.heading_size(b.level))

    elif b.kind == "para":
        p = doc.add_paragraph()
        pf = p.paragraph_format
        pf.first_line_indent = Pt(0)
        _emit_runs(p, b.runs, base_dir)

    elif b.kind == "code":
        _render_code(doc, b)

    elif b.kind == "quote":
        _render_quote(doc, b, base_dir)

    elif b.kind == "list":
        _render_list(doc, b, base_dir, level=ctx.get("level", 1))

    elif b.kind == "table":
        _render_table(doc, b, base_dir)

    elif b.kind == "hr":
        _render_hr(doc)

    elif b.kind == "image":
        _render_image(doc, b.src, b.alt, base_dir)


def _emit_runs(p, runs: list[Run], base_dir: str, **base_style):
    for r in runs:
        if r.brk:
            p.add_run().add_break()
            continue
        if r.image:
            _render_image(p, r.image, r.text, base_dir)
            continue
        if not r.text:
            continue
        if r.link:
            add_hyperlink(p, r.link, r.text)
            continue
        run = p.add_run(r.text)
        style_run(run,
                  name=theme.BODY_FONT, ea=theme.BODY_FONT_EA,
                  size=base_style.get("size", theme.BODY_SIZE),
                  bold=base_style.get("bold", r.bold) or None,
                  italic=r.italic or None,
                  strike=r.strike or None,
                  sup=r.sup, sub=r.sub,
                  mono=r.code,
                  color=theme.CODE_FG if r.code else base_style.get("color"),
                  shading=theme.CODE_BG if r.code else None)


def _render_code(doc, b: Block):
    lines = b.text.split("\n") or [""]
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(6)
    pf.space_after = Pt(6)
    pf.left_indent = Cm(0.2)
    pf.right_indent = Cm(0.2)
    pf.line_spacing = 1.0
    _shd(p._p.get_or_add_pPr(), theme.CODE_BG)
    for i, line in enumerate(lines):
        if i:
            p.add_run().add_break()
        run = p.add_run(line if line else " ")
        style_run(run, mono=True, size=9.5, color=theme.CODE_FG)


def _render_quote(doc, b: Block, base_dir: str, depth: int = 1):
    for child in b.items:
        if child.kind == "quote":
            _render_quote(doc, child, base_dir, depth + 1)
            continue
        if child.kind in ("list", "table", "code", "hr", "image"):
            _render_block(doc, child, base_dir)
            continue
        p = doc.add_paragraph()
        pf = p.paragraph_format
        pf.left_indent = Cm(0.6 * depth)
        pf.space_after = Pt(4)
        # 左侧竖线
        pPr = p._p.get_or_add_pPr()
        borders = OxmlElement("w:pBdr")
        left = _sub(borders, "w:left", val="single", sz="18", space="8",
                    color=theme.QUOTE_BAR)
        pPr.append(borders)
        _shd(pPr, theme.QUOTE_BG)
        for r in (child.runs or []):
            if r.brk:
                p.add_run().add_break()
                continue
            if r.image:
                _render_image(p, r.image, r.text, base_dir)
                continue
            if not r.text:
                continue
            if r.link:
                add_hyperlink(p, r.link, r.text)
                continue
            run = p.add_run(r.text)
            style_run(run, ea=theme.BODY_FONT_EA, size=10.5, italic=True,
                      bold=r.bold or None, strike=r.strike or None,
                      color=theme.LINK if r.link else theme.QUOTE_FG,
                      mono=r.code,
                      shading=theme.CODE_BG if r.code else None)


def _render_list(doc, b: Block, base_dir: str, level: int = 1):
    style_base = "List Number" if b.ordered else "List Bullet"
    style_name = style_base
    if level > 1:
        cand = f"{style_base} {min(level, 3)}"
        try:
            doc.styles[cand]
            style_name = cand
        except KeyError:
            style_name = style_base

    for idx, item in enumerate(b.items, start=b.start):
        first = True
        for sub in item:
            if sub.kind == "list":
                _render_list(doc, sub, base_dir, level + 1)
                continue
            if sub.kind == "table":
                _render_table(doc, sub, base_dir)
                continue
            if sub.kind == "code":
                _render_code(doc, sub)
                continue
            try:
                p = doc.add_paragraph(style=style_name)
            except KeyError:
                p = doc.add_paragraph()
            pf = p.paragraph_format
            pf.space_after = Pt(2)
            pf.line_spacing = 1.4
            if style_name == style_base and level > 1:
                pf.left_indent = Cm(0.75 * level)
                pf.first_line_indent = Cm(-0.4)
            if not first:
                pf.space_before = Pt(0)
            _emit_runs(p, sub.runs or [], base_dir)
            first = False


def _render_table(doc, b: Block, base_dir: str):
    rows = b.rows
    if not rows:
        return
    n_cols = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=n_cols)
    try:
        table.style = "Table Grid"
    except KeyError:
        pass
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True

    for ri, row in enumerate(rows):
        is_head = b.header and ri == 0
        for ci in range(n_cols):
            cell = table.cell(ri, ci)
            cell.width = Cm(16.0 / n_cols)
            blocks = row[ci] if ci < len(row) else []
            first = True
            for blk in blocks:
                if first:
                    p = cell.paragraphs[0]
                    first = False
                else:
                    p = cell.add_paragraph()
                pf = p.paragraph_format
                pf.space_after = Pt(2)
                pf.line_spacing = 1.2
                runs = blk.runs or []
                for r in runs:
                    if r.brk:
                        p.add_run().add_break()
                        continue
                    if r.image:
                        _render_image(p, r.image, r.text, base_dir)
                        continue
                    if not r.text:
                        continue
                    if r.link:
                        add_hyperlink(p, r.link, r.text)
                        continue
                    run = p.add_run(r.text)
                    style_run(run, ea=theme.BODY_FONT_EA, size=10,
                              bold=(is_head or r.bold) or None,
                              italic=r.italic or None,
                              strike=r.strike or None,
                              mono=r.code,
                              color=theme.CODE_FG if r.code else None,
                              shading=theme.CODE_BG if r.code else None)
            # 表头底色
            tcPr = cell._tc.get_or_add_tcPr()
            if is_head:
                _shd(tcPr, theme.TABLE_HEAD_BG)
            elif ri % 2 == 1:
                _shd(tcPr, theme.TABLE_ALT_BG)

    if b.header and rows:
        trPr = table.rows[0]._tr.get_or_add_trPr()
        _sub(trPr, "w:tblHeader", val="true")

    doc.add_paragraph()


def _render_hr(doc):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(6)
    pf.space_after = Pt(6)
    pPr = p._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    _sub(borders, "w:bottom", val="single", sz="6", space="1", color=theme.HR)
    pPr.append(borders)


def _render_image(target, src: str, alt: str, base_dir: str):
    """target 为 Document 时独占一段并居中，为 Paragraph 时行内插入。"""
    own_para = hasattr(target, "add_paragraph")
    para = target.add_paragraph() if own_para else target

    def placeholder():
        run = para.add_run(alt or "[图片]")
        style_run(run, ea=theme.BODY_FONT_EA, size=10, italic=True,
                  color=theme.QUOTE_FG)

    got = load_image(src, base_dir)
    if not got:
        placeholder()
        return
    data, _ext = got
    try:
        shape = para.add_run().add_picture(io.BytesIO(data))
    except Exception:
        placeholder()
        return
    try:
        doc = para.part.document
        max_w = _content_width(doc)
        if shape.width > max_w:
            ratio = shape.height / shape.width
            shape.width = int(max_w)
            shape.height = int(max_w * ratio)
    except Exception:
        pass
    if own_para:
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
